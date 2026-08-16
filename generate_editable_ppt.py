#!/usr/bin/env python3
"""Generate a fully editable PowerPoint from the evidence-backed HTML report."""

from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT / "diagrams"))
from generate_evidence_diagrams import (  # noqa: E402
    CATEGORY,
    spec_databricks,
    spec_oceanbase,
    spec_oracle,
    spec_polardb,
)
from render_final_html import SOURCES  # noqa: E402


SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Microsoft YaHei"

NAVY = "0F172A"
SLATE = "475569"
LIGHT = "F8FAFC"
LINE = "DBE4EF"
BLUE = "2563EB"
GREEN = "16A34A"
CYAN = "0891B2"
ORANGE = "EA580C"
GRAY = "64748B"
WHITE = "FFFFFF"


def rgb(value):
    return RGBColor.from_string(value.replace("#", ""))


def set_run(run, size=16, bold=False, color=NAVY, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(slide, x, y, w, h, text, size=16, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = text
    set_run(r, size, bold, color)
    return shape


def add_rich_lines(slide, x, y, w, h, lines, size=15, color=NAVY,
                   bullet=True, level=0, spacing=5):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = level
        p.space_after = Pt(spacing)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        if bullet:
            p.text = "• " + line
    return shape


def add_box(slide, x, y, w, h, title, body="", fill="EFF6FF",
            stroke=BLUE, title_size=15, body_size=11):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = Pt(1.4)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    set_run(r, title_size, True, stroke)
    if body:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(4)
        r2 = p2.add_run()
        r2.text = body
        set_run(r2, body_size, False, SLATE)
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=SLATE, width=1.2):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    conn.line.dash_style = MSO_LINE.SOLID
    ln = conn._element.spPr.ln
    head = OxmlElement("a:headEnd")
    head.set("type", "triangle")
    ln.append(head)
    return conn


def add_title(slide, title, subtitle=None, section=None):
    add_text(slide, 0.55, 0.28, 12.2, 0.45, title, 24, NAVY, True)
    if subtitle:
        add_text(slide, 0.58, 0.78, 11.9, 0.30, subtitle, 10.5, SLATE)
    if section:
        add_text(slide, 11.7, 0.30, 1.0, 0.28, section, 10, BLUE, True,
                 PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.08), Inches(12.2), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(LINE)
    line.line.fill.background()


def add_footer(slide, page, source_text="官方资料编号见末页"):
    add_text(slide, 0.55, 7.12, 10.7, 0.20, source_text, 8.5, GRAY)
    add_text(slide, 12.0, 7.10, 0.7, 0.20, str(page), 9, GRAY, False,
             PP_ALIGN.RIGHT)


def new_slide(prs, title=None, subtitle=None, section=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(WHITE)
    if title:
        add_title(slide, title, subtitle, section)
    return slide


def add_table(slide, x, y, w, h, headers, rows, widths=None,
              font_size=10.5):
    table = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    for c, value in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = value
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                set_run(r, font_size, True, WHITE)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("F8FAFC" if r_idx % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    set_run(run, font_size, c_idx == 0, NAVY if c_idx == 0 else SLATE)
    return table


def add_diagram(slide, spec, x=0.45, y=1.25, w=12.45, h=5.75):
    """Render spec as editable native PowerPoint shapes and connectors."""
    nodes = {n[0]: n for n in spec["nodes"]}
    sx, sy = w / 1120.0, h / 620.0

    # Connectors first so editable boxes remain above lines.
    for src, dst, label in spec["edges"]:
        _, ax, ay, aw, ah, _, _ = nodes[src]
        _, bx, by, bw, bh, _, _ = nodes[dst]
        if abs(by - ay) < 120:
            x1, y1 = x + (ax + aw) * sx, y + (ay + ah / 2) * sy
            x2, y2 = x + bx * sx, y + (by + bh / 2) * sy
        else:
            x1, y1 = x + (ax + aw / 2) * sx, y + (ay + ah) * sy
            x2, y2 = x + (bx + bw / 2) * sx, y + by * sy
        add_arrow(slide, x1, y1, x2, y2, SLATE, 1.0)
        add_text(slide, (x1 + x2) / 2 - 0.7, (y1 + y2) / 2 - 0.16,
                 1.4, 0.28, label, 6.5, SLATE, False, PP_ALIGN.CENTER)

    for _, nx, ny, nw, nh, cat, label in spec["nodes"]:
        cat_label, fill, stroke = CATEGORY[cat]
        px, py, pw, ph = x + nx * sx, y + ny * sy, nw * sx, nh * sy
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(py),
            Inches(pw), Inches(ph)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(stroke)
        shape.line.width = Pt(1.4)
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = Inches(0.035)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = cat_label
        set_run(r, 6.5, True, stroke)
        for idx, line in enumerate(label.splitlines()):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            p2.space_before = p2.space_after = Pt(0)
            r2 = p2.add_run()
            r2.text = line
            set_run(r2, 8 if idx == 0 else 6.8, idx == 0, stroke if idx == 0 else SLATE)

    # Editable legend.
    lx = x + 0.1
    for cat in ("kernel", "control", "sync", "service", "storage"):
        label, fill, stroke = CATEGORY[cat]
        add_box(slide, lx, y + h - 0.28, 2.18, 0.22, label, "", fill, stroke, 7.2)
        lx += 2.35


def cover(prs):
    slide = new_slide(prs)
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(NAVY)
    add_text(slide, 0.75, 1.15, 11.8, 1.4,
             "AI 时代数据库内核增强、\n库仓一体与 Agent 架构",
             31, WHITE, True, PP_ALIGN.LEFT)
    add_text(slide, 0.78, 2.85, 10.8, 0.7,
             "Oracle 23ai/26ai · PolarDB · OceanBase · Databricks\n"
             "基于官方资料的技术架构、组件落点与客户实践",
             16, "CBD5E1")
    add_text(slide, 0.78, 6.35, 5.2, 0.35,
             "可编辑 PowerPoint · 2026-08-16", 11, "94A3B8")


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    cover(prs)
    page = 2

    slide = new_slide(prs, "研究原则与报告结构", "只采用官方文档、官方产品页、官方博客和官方客户案例", "方法")
    add_box(slide, 0.65, 1.35, 3.8, 1.3, "厂商事实", "每个架构组件和技术结论均标注 [O/P/B/D] 来源编号", "EFF6FF", BLUE)
    add_box(slide, 4.75, 1.35, 3.8, 1.3, "证据边界", "公开资料未披露的进程拓扑、事务协议和部署形态明确标为“未披露”", "F0FDF4", GREEN)
    add_box(slide, 8.85, 1.35, 3.8, 1.3, "自研建议", "与厂商事实分离：说明哪些应进内核、复制层、AI 节点和控制面", "FFF7ED", ORANGE)
    add_rich_lines(slide, 0.8, 3.05, 11.7, 2.7, [
        "Oracle：数据库双格式、Vector/ONNX、Select AI Agent、托管 MCP、Iceberg Lakehouse",
        "PolarDB：InnoDB IMCI、Redo/物理复制、专属 AI 节点与 Serverless Data-Agent",
        "OceanBase：HTAP/向量/AI 函数、External Catalog、Lakebase 产品路线及披露边界",
        "Databricks：Lakebase 存储、LTAP 转码、AI Search 派生索引、Agent Bricks 与 Unity Catalog",
    ], 16)
    add_footer(slide, page)
    page += 1

    slide = new_slide(prs, "四家厂商的共同终局与不同路径", "都在利用既有数据，但不是同一种实现", "总览")
    headers = ["厂商", "既有数据", "补 AP/仓", "补 AI/Agent", "路径本质"]
    rows = [
        ["Oracle", "Oracle 行格式业务数据、外部 Iceberg", "实例 SGA 内双格式 IM Column Store", "Vector/ONNX/Select AI Agent/托管 MCP", "单一数据库内核增强"],
        ["PolarDB", "InnoDB 行存与共享存储", "InnoDB IMCI 二级索引；RO 节点 Redo 维护", "专属 AI 节点或 Serverless；SQL MLOps", "内核 HTAP + 独立 AI 算力"],
        ["OceanBase", "OceanBase 事务数据与外部湖数据", "HTAP；External Catalog；Lakebase 路线", "向量/语义索引/AI 函数；上层 Agent 产品", "数据库向多模和开放湖扩展"],
        ["Databricks", "Delta/Iceberg 与 Lakebase Postgres", "LTAP 存储转码为 Parquet", "AI Search 与 Agent Bricks 托管服务", "湖仓向 OLTP 扩展"],
    ]
    add_table(slide, 0.55, 1.35, 12.2, 4.75, headers, rows,
              [1.1, 2.4, 2.9, 3.1, 2.7], 9.5)
    add_text(slide, 0.65, 6.35, 11.8, 0.5,
             "关键判断：数据库内核负责确定性数据处理；模型算力和 Agent 编排是否进内核，四家选择不同。",
             15, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide, page)
    page += 1

    # Oracle architecture.
    slide = new_slide(prs, "Oracle 26ai：数据库内核扩展 AI 与湖仓", "26ai 是 23ai 的后续长期支持版本 [O0]", "Oracle")
    add_diagram(slide, spec_oracle())
    add_footer(slide, page, "来源：[O0]-[O7]")
    page += 1

    slide = new_slide(prs, "Oracle：TP、AP、AI 与湖如何协同", section="Oracle")
    add_box(slide, 0.6, 1.3, 3.9, 1.45, "① 双格式内核", "Buffer Cache 保存行格式，IM Column Store 保存列格式；优化器透明选择，数据库后台进程填充/重填 [O1][O2]", "EFF6FF", BLUE)
    add_box(slide, 4.72, 1.3, 3.9, 1.45, "② Converged AI", "VECTOR、关系、JSON、Graph 在同一数据库；可加载本地 ONNX，也可调用外部模型 [O1][O6]", "EFF6FF", BLUE)
    add_box(slide, 8.84, 1.3, 3.9, 1.45, "③ Agent 与 MCP", "Select AI Agent 由 DBMS_CLOUD_AI_AGENT 管理；Autonomous DB 提供每数据库托管 MCP endpoint [O4][O5]", "F0FDF4", GREEN)
    add_box(slide, 0.6, 3.2, 5.9, 1.55, "④ Exadata 下推", "AI Smart Scan 在存储服务器执行向量距离计算和 Top-K，减少网络传输和 DB Server 处理 [O3]", "F8FAFC", GRAY)
    add_box(slide, 6.75, 3.2, 5.95, 1.55, "⑤ Iceberg 湖仓", "Autonomous AI Lakehouse 原位查询 Apache Iceberg，并使用统一 Catalog 与 Data Lake Accelerator [O7]", "F8FAFC", GRAY)
    add_rich_lines(slide, 0.8, 5.25, 11.8, 1.25, [
        "纠正：IM Column Store 不是外部 CDC/Redo 直灌服务；它位于数据库实例 SGA。",
        "纠正：Autonomous AI Database MCP 是托管能力，不要求客户自建 Sidecar。",
    ], 14, ORANGE)
    add_footer(slide, page, "来源：[O1]-[O7]")
    page += 1

    slide = new_slide(prs, "Oracle 官方客户实践", "利用原有数据库数据直接叠加 Vector Search 与 Select AI", "Oracle")
    add_box(slide, 0.8, 1.45, 5.75, 3.7, "Retraced", "供应商生命周期与合规平台\n\n在原有 Autonomous Database 上升级 26ai，使用 AI Vector Search 和 Select AI。\n\n官方结果：供应商重复记录减少 80%。\n\n来源：[O8]", "EFF6FF", BLUE, 20, 15)
    add_box(slide, 6.8, 1.45, 5.75, 3.7, "Rappi", "拉美即时配送平台\n\nAutonomous AI Database + AI Vector Search 直接处理既有商品目录，避免在专用数据库间搬数据。\n\n官方结果：搜索响应延迟降低 40%。\n\n来源：[O9]", "F0FDF4", GREEN, 20, 15)
    add_text(slide, 1.0, 5.55, 11.2, 0.75,
             "客户价值：既有交易/目录数据不必先迁移到独立向量库，即可增加语义检索和自然语言访问。",
             18, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, page, "来源：[O8][O9]")
    page += 1

    # PolarDB.
    slide = new_slide(prs, "PolarDB：InnoDB IMCI + 独立 AI 节点", "内核 HTAP 与 AI 算力解耦", "PolarDB")
    add_diagram(slide, spec_polardb())
    add_footer(slide, page, "来源：[P1]-[P8]")
    page += 1

    slide = new_slide(prs, "PolarDB：内核与服务部署边界", section="PolarDB")
    headers = ["组件", "公开实现", "物理落点", "为什么这样设计"]
    rows = [
        ["IMCI", "InnoDB 列存二级索引；默认驻内存，可落共享存储", "数据库内核 / RO AP 节点", "复用事务、Redo、物理复制和 MySQL 兼容性"],
        ["Redo 同步", "RW 生成 Redo；RO/Standby 后台构建和维护列索引", "数据库复制层", "避免 Kafka/Flink 外部 ETL，降低主节点写入影响"],
        ["优化器", "行存串行、行存并行、IMCI 三路径成本选择", "数据库内核", "同时支持 TP 与 AP，查询自动路由"],
        ["PolarDB for AI", "SQL MLOps、NL2SQL、模型训练/推理", "专属 AI Node 或 Serverless", "隔离 GPU/AI 负载，支持弹性与独立计费"],
        ["Data-Agent", "NL2SQL、NL2Chart、Summary", "AI 服务/控制面入口", "让业务用户直接分析既有数据库数据"],
    ]
    add_table(slide, 0.55, 1.3, 12.2, 4.95, headers, rows,
              [1.55, 3.8, 2.55, 4.3], 10)
    add_text(slide, 0.65, 6.45, 11.9, 0.4,
             "准确表述：PolarDB AI 节点是官方独立节点或 Serverless 资源，不能笼统称为“同机 Sidecar”。",
             14, ORANGE, True, PP_ALIGN.CENTER)
    add_footer(slide, page, "来源：[P1]-[P6]")
    page += 1

    slide = new_slide(prs, "PolarDB 官方客户：雅迪", "DMS + PolarDB for AI 将自然语言分析带到销售数据", "PolarDB")
    add_box(slide, 0.7, 1.35, 3.8, 1.45, "既有数据", "批发、销售、库存、采购、供应链与营销数据保留在 PolarDB", "EFF6FF", BLUE)
    add_box(slide, 4.78, 1.35, 3.8, 1.45, "能力组合", "DMS + PolarDB for AI\nNL2SQL + RAG + 表格/图表/文字解读", "F0FDF4", GREEN)
    add_box(slide, 8.86, 1.35, 3.8, 1.45, "业务入口", "雅迪云销通 App\n销售人员以口语提问直接获取经营分析", "FFF7ED", ORANGE)
    add_arrow(slide, 4.5, 2.05, 4.78, 2.05, BLUE, 2)
    add_arrow(slide, 8.58, 2.05, 8.86, 2.05, BLUE, 2)
    add_text(slide, 0.8, 3.35, 11.7, 0.55, "官方披露结果", 20, NAVY, True, PP_ALIGN.CENTER)
    add_box(slide, 1.0, 4.0, 3.35, 1.45, "10 万+", "销售与经销商人员使用", "EFF6FF", BLUE, 26, 15)
    add_box(slide, 4.98, 4.0, 3.35, 1.45, "90%+", "NL2SQL 查询准确率", "F0FDF4", GREEN, 26, 15)
    add_box(slide, 8.96, 4.0, 3.35, 1.45, "实时", "获取全域业务数据分析", "FFF7ED", ORANGE, 26, 15)
    add_footer(slide, page, "来源：[P8]")
    page += 1

    # OceanBase.
    slide = new_slide(prs, "OceanBase：数据库多模能力 + 外部湖接入", "区分已发布能力与 Lakebase 产品路线", "OceanBase")
    add_diagram(slide, spec_oceanbase())
    add_footer(slide, page, "来源：[B1]-[B7]")
    page += 1

    slide = new_slide(prs, "OceanBase：哪些已有文档，哪些仍是产品路线", section="OceanBase")
    headers = ["能力", "证据状态", "当前可确认实现", "不能臆测的部分"]
    rows = [
        ["AI 函数", "官方产品文档", "AI_SPLIT_DOCUMENT / EMBED / COMPLETE / RERANK 作为 SQL 表达式；DBMS_AI_SERVICE 注册 endpoint", "模型服务内部部署由 endpoint 决定，不能一概标为内核模型"],
        ["语义索引", "官方产品文档", "VARCHAR 自动 embedding 并建向量索引；查询可输入原始文本", "具体异步队列与事务协议未完整公开"],
        ["External Catalog", "官方产品文档", "HMS / REST / FILESYSTEM / ODPS；读取 Iceberg snapshot/manifest", "当前以只读为主，不能写成已实现完整双向共享写"],
        ["Spark 集成", "官方最佳实践", "Spark Catalog / Connector 映射元数据、并行读写、谓词下推", "不等于 Spark 直接读取 OceanBase 内核页格式"],
        ["Lakebase/Fork/AI列", "官方方案页", "S3/Iceberg/统一 Catalog/Spark/Ray、多模表、Fork 等产品目标", "进程拓扑、跨引擎写冲突和事务协议未完全披露"],
    ]
    add_table(slide, 0.45, 1.25, 12.45, 5.25, headers, rows,
              [1.55, 1.45, 4.3, 5.15], 9.4)
    add_footer(slide, page, "来源：[B1]-[B7]")
    page += 1

    slide = new_slide(prs, "OceanBase：利用既有数据构造 AI 数据底座", section="OceanBase")
    add_box(slide, 0.65, 1.35, 3.8, 1.4, "数据库内核", "事务、HTAP、关系过滤、全文、向量与语义索引\n同一 SQL 查询路径 [B1][B7]", "EFF6FF", BLUE)
    add_box(slide, 4.75, 1.35, 3.8, 1.4, "模型服务", "AI 函数在 SQL 中；通过 DBMS_AI_SERVICE 注册第三方模型 endpoint [B2][B3]", "FFF7ED", ORANGE)
    add_box(slide, 8.85, 1.35, 3.8, 1.4, "外部湖", "External Catalog 读取 OSS/S3/HDFS 中 Iceberg/Parquet，当前以只读查询为主 [B4][B5]", "F8FAFC", GRAY)
    add_box(slide, 1.2, 3.35, 5.25, 1.55, "产品与控制层", "Fork Database、Agent 沙箱、上下文层、PowerMem、DataPilot\n官方方案页确认能力方向，但内部拓扑尚未完整披露 [B1]", "F0FDF4", GREEN)
    add_box(slide, 6.85, 3.35, 5.25, 1.55, "开放计算", "Spark Catalog/Connector 已有官方实践；Lakebase 方案页进一步提出 Spark/Ray 对统一数据工作 [B1][B6]", "F0FDF4", GREEN)
    add_text(slide, 0.8, 5.55, 11.7, 0.75,
             "对自研团队的启示：先把可验证的内核向量/全文/HTAP和 External Catalog 做实，再逐步扩展多模表、AI 列与跨引擎写。",
             17, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, page, "来源：[B1]-[B7]")
    page += 1

    # Databricks.
    slide = new_slide(prs, "Databricks：Lakebase LTAP + Agent Bricks", "从湖仓向 OLTP 扩展；数据与 Agent 分层", "Databricks")
    add_diagram(slide, spec_databricks())
    add_footer(slide, page, "来源：[D1]-[D8]")
    page += 1

    slide = new_slide(prs, "Databricks：必须区分三条路径", section="Databricks")
    add_box(slide, 0.55, 1.35, 3.95, 4.55, "① LTAP 数据路径",
            "Lakebase Postgres Compute\n↓\nSafekeepers / Pageservers / Object Storage\n↓\n存储物化时转码为 Parquet\n↓\nDelta / Iceberg\n↓\nLakehouse//RT 直接分析\n\n不是外部 CDC/复制管道\n[D1][D2][D3]",
            "EFF6FF", BLUE, 19, 14)
    add_box(slide, 4.7, 1.35, 3.95, 4.55, "② AI Search 路径",
            "Delta Table\n↓\nDelta Sync Index\n↓\n托管 Serverless AI Search\n↓\nHNSW / BM25 / Hybrid Search\n\n这是派生搜索索引，不等于 LTAP 单一逻辑副本\n[D6]",
            "FFF7ED", ORANGE, 19, 14)
    add_box(slide, 8.85, 1.35, 3.95, 4.55, "③ Agent 路径",
            "Supervisor Agent\n↓\nGenie / AI Search / UC Functions\nMCP / Custom Agents\n↓\nDatabricks Apps / Model Serving\n\n由 Unity Catalog 和 AI Gateway 统一治理\n[D4][D5][D7]",
            "F0FDF4", GREEN, 19, 14)
    add_footer(slide, page, "来源：[D1]-[D7]")
    page += 1

    slide = new_slide(prs, "Databricks 官方客户：FinThrive", "在既有平台、Notebook、代码和文档上快速构建 Agent", "Databricks")
    add_rich_lines(slide, 0.8, 1.35, 5.65, 4.7, [
        "既有资产：数百个 Databricks Notebooks、遗留数据管道、Wiki 与内部文档",
        "Unity Catalog：将 Notebook、代码库和文档以受治理资产暴露",
        "Agent Bricks：分别构建 Databricks 代码、Azure Data Factory 和内部文档 Agent",
        "Supervisor：把多个 Agent 的答案统一编排",
        "Databricks App：向员工提供聊天界面",
    ], 15)
    add_box(slide, 7.0, 1.45, 5.1, 3.1, "官方案例结果",
            "“Within a few hours”\n数小时内上线可查询数据管道的助手\n\n员工能够以自然语言获得带引用的代码与管道答案\n\n来源：[D8]",
            "EFF6FF", BLUE, 22, 16)
    add_text(slide, 7.1, 5.0, 4.9, 0.9,
             "关键价值：不是先迁移数据，\n而是在既有 Databricks 数据和治理基础上增加 Agent。",
             16, GREEN, True, PP_ALIGN.CENTER)
    add_footer(slide, page, "来源：[D8]")
    page += 1

    # Self-developed mapping.
    slide = new_slide(prs, "自研数据库：组件落点与研发边界", "以数据库为核心，但不把所有 AI 逻辑塞进事务进程", "建议")
    headers = ["目标", "建议落点", "研发工作", "参考实现"]
    rows = [
        ["既有 TP 数据直接做 AP", "数据库内核", "列存格式/列索引、向量化执行器、CBO 行列路径选择、资源隔离", "Oracle IMCS / PolarDB IMCI"],
        ["保持行列实时一致", "内核后台进程或复制层", "后台填充/重填，或 Redo/物理复制到 RO 列存节点", "Oracle / PolarDB"],
        ["统一业务查询与语义检索", "数据库内核", "VECTOR 类型、ANN 索引、全文、混合优化器、RLS", "Oracle / OceanBase"],
        ["Embedding / LLM", "库内轻量模型 + 独立 AI Node", "轻量 ONNX 可库内；GPU/大模型放专属节点或 Serverless", "Oracle / PolarDB"],
        ["NL2SQL 与 Agent", "托管服务/控制面", "Agent 编排、模型路由、评测、Tracing、权限传递", "四家厂商"],
        ["湖数据接入", "Catalog + Connector", "Iceberg REST/HMS、Parquet Reader、谓词下推、凭据与权限", "Oracle / OB / Databricks"],
    ]
    add_table(slide, 0.45, 1.25, 12.45, 5.55, headers, rows,
              [2.0, 2.15, 5.0, 3.3], 9.8)
    add_footer(slide, page, "建议基于官方实现对比")
    page += 1

    slide = new_slide(prs, "一键生成知识图谱：端到端产品流程",
                      "“一键”是自动化工作流，不取消本体与实体冲突审核", "知识图谱")
    stages = [
        ("1 数据接入", "表/文档发现\n一致性快照\nLSN 水位", BLUE),
        ("2 本体生成", "实体/关系/属性\n业务术语\n人工审核", GREEN),
        ("3 实体关系抽取", "SQL 确定性抽取\nNER/RE\n证据片段", ORANGE),
        ("4 实体消歧", "主键/规则\n向量相似\n邻居匹配", CYAN),
        ("5 图谱持久化", "实体/边/属性\n时态版本\n来源血缘", BLUE),
        ("6 增量维护", "Redo/CDC\n删除传播\n模型重算", CYAN),
        ("7 服务发布", "图查询\nGraphRAG\nSQL/REST/MCP", GREEN),
    ]
    x, y = 0.35, 1.45
    for idx, (title, body, color) in enumerate(stages):
        add_box(slide, x, y, 1.65, 3.5, title, body, "FFFFFF", color, 13.5, 11.5)
        if idx < len(stages) - 1:
            add_arrow(slide, x + 1.67, y + 1.75, x + 1.82, y + 1.75, SLATE, 1.6)
        x += 1.82
    add_text(slide, 0.65, 5.35, 12.0, 0.55,
             "结构化关系优先用 SQL/Join 确定性生成；LLM 只补充文档中的非结构化实体和关系。",
             16, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, 0.65, 6.05, 12.0, 0.55,
             "每个实体和关系必须保留 source_pk、source_lsn、ontology/model 版本、confidence 与 evidence。",
             14, ORANGE, True, PP_ALIGN.CENTER)
    add_footer(slide, page, "建议架构；GraphRAG 参考：[D9]")
    page += 1

    slide = new_slide(prs, "一键知识图谱：技术能力与工程落点", section="知识图谱")
    headers = ["技术域", "必须能力", "产品组件", "工程落点"]
    rows = [
        ["一致性基础", "Snapshot、Redo/CDC、水位、Schema 演进、删除传播", "Graph Snapshot / Change Feed", "TP 内核 + 复制层"],
        ["本体语义", "实体/关系/属性约束、版本、审批、兼容迁移", "Ontology Studio", "管控面"],
        ["结构化抽取", "PK/FK 发现、SQL/Join、规则、时态关系", "Deterministic Graph Builder", "AP 节点"],
        ["非结构化抽取", "OCR、Parser、Chunk、NER、关系抽取、证据定位", "Document Intelligence Worker", "AI Worker"],
        ["实体解析", "规范化、规则、向量、图邻居匹配、人工审核", "Entity Resolution Service", "AP + AI 服务"],
        ["图查询", "实体/边、图索引、多跳、时态图、社区与中心性", "Graph Tables / Property Graph", "数据库/AP"],
        ["混合召回", "关系、全文、向量、图邻居、Rerank", "GraphRAG Retriever", "数据库 + AI Search"],
        ["治理质量", "RLS、血缘、置信度、模型/Prompt 版本、评测", "Graph Quality Center", "内核安全 + 管控面"],
    ]
    add_table(slide, 0.42, 1.22, 12.5, 5.78, headers, rows,
              [1.5, 4.15, 3.65, 3.2], 8.9)
    add_footer(slide, page)
    page += 1

    slide = new_slide(prs, "这套架构可以对外提供什么解决方案", section="解决方案")
    headers = ["解决方案", "核心产品组合", "目标行业", "客户价值"]
    rows = [
        ["实时经营分析助手", "TP + AP + 自动同步 + Semantic View + ChatBI", "零售、制造、物流、SaaS", "最新业务数据直接分析，减少 T+1 数仓和人工报表"],
        ["企业知识与业务问答", "业务表 + 文档 + AI Index + 混合检索 + RAG", "金融、政企、运营商、制造", "制度知识与实时业务状态统一回答"],
        ["供应链知识图谱", "ERP 表 + 合同/质检文档 + 消歧 + GraphRAG", "制造、汽车、零售", "供应商穿透、替代关系、质量和交付风险"],
        ["客户 360 与智能营销", "客户/订单/行为图谱 + 实时特征 + 推荐 Agent", "零售、金融、互联网", "统一客户身份、下一最佳行动与流失预测"],
        ["反欺诈关系网络", "实时交易 + 设备/账户图谱 + 图算法 + 风险模型", "银行、支付、保险", "团伙识别、风险传播和可解释证据链"],
        ["设备运维知识图谱", "设备台账 + 告警/工单 + 拓扑 + 根因 Agent", "能源、运营商、工业", "故障定位、影响分析和维修建议"],
        ["Agent 数据底座", "MCP + Memory + 任务状态 + SQL Tool + 沙箱审计", "AI SaaS、企业 AI 中台", "Agent 安全读取、分析和修改业务状态"],
    ]
    add_table(slide, 0.42, 1.22, 12.5, 5.78, headers, rows,
              [2.15, 4.25, 2.5, 3.6], 9.0)
    add_footer(slide, page)
    page += 1

    slide = new_slide(prs, "建议的产品演进路线", "先内核确定性能力，再知识库/ChatBI，再知识图谱与 Agent", "建议")
    stages = [
        ("阶段 1\nHTAP", "AP 节点\nRedo 增量维护\nCBO 行列选择\n资源隔离", BLUE),
        ("阶段 2\nAI 数据", "VECTOR/全文\nAI Gateway\n异步 AI Job\n混合检索", CYAN),
        ("阶段 3\n知识与 BI", "文档智能\nAI Index\nSemantic View\nNL2SQL/RAG", GREEN),
        ("阶段 4\n知识图谱", "Ontology\n抽取与消歧\n增量图谱\nGraphRAG", ORANGE),
        ("阶段 5\nAgent Native", "MCP/Memory\n任务状态\n数据沙箱\nTool 审计", BLUE),
    ]
    x = 0.3
    for idx, (title, body, color) in enumerate(stages):
        add_box(slide, x, 1.55, 2.35, 3.7, title, body, "FFFFFF", color, 16, 13)
        if idx < len(stages) - 1:
            add_arrow(slide, x + 2.37, 3.35, x + 2.52, 3.35, SLATE, 1.8)
        x += 2.58
    add_text(slide, 0.8, 5.75, 11.7, 0.65,
             "原则：每一阶段都直接利用当前数据库中的既有数据；避免先建设一套独立大数据平台再回接数据库。",
             17, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, page)
    page += 1

    slide = new_slide(prs, "目标客户与解决方案", section="商业")
    add_box(slide, 0.65, 1.3, 3.8, 2.0, "金融 / 政企 / 运营商",
            "核心诉求\n数据不出域、强一致、审计、实时风控\n\n产品组合\nTP + 实时 AP + 向量/全文 + AI Node", "EFF6FF", BLUE, 18, 13)
    add_box(slide, 4.75, 1.3, 3.8, 2.0, "零售 / 制造 / 物流",
            "核心诉求\n实时经营分析、自然语言问数、知识助手\n\n产品组合\nIMCI/列存 + Data-Agent + RAG", "F0FDF4", GREEN, 18, 13)
    add_box(slide, 8.85, 1.3, 3.8, 2.0, "互联网 / AI SaaS",
            "核心诉求\n搜索推荐、Agent 状态、低成本弹性\n\n产品组合\n混合搜索 + 分支沙箱 + Serverless AI", "FFF7ED", ORANGE, 18, 13)
    add_text(slide, 0.8, 3.85, 11.7, 0.45, "客户购买的不是“AI 函数”，而是更短的数据链路与更低的一致性、运维和治理成本。", 18, NAVY, True, PP_ALIGN.CENTER)
    add_rich_lines(slide, 1.1, 4.55, 11.1, 1.5, [
        "替代或减少：业务库 → Kafka/CDC → 数仓 → 搜索/向量库 → Agent 的多套同步链路",
        "保留优势：现有 SQL、事务、权限、数据模型和客户业务连续性",
        "新增价值：实时分析、语义搜索、自然语言问数和可治理 Agent",
    ], 15)
    add_footer(slide, page)
    page += 1

    # References, split into 4 slides.
    groups = [
        ("Oracle 官方资料", [k for k in SOURCES if k.startswith("O")]),
        ("PolarDB 官方资料", [k for k in SOURCES if k.startswith("P")]),
        ("OceanBase 官方资料", [k for k in SOURCES if k.startswith("B")]),
        ("Databricks 官方资料", [k for k in SOURCES if k.startswith("D")]),
    ]
    for group_title, keys in groups:
        slide = new_slide(prs, group_title, "所有 URL 均为厂商官方文档、官方产品页、官方博客或官方客户案例", "参考")
        y = 1.3
        for key in keys:
            title, url = SOURCES[key]
            add_text(slide, 0.7, y, 1.0, 0.28, f"[{key}]", 12, BLUE, True)
            shape = add_text(slide, 1.45, y, 4.4, 0.30, title, 11.5, NAVY, True)
            add_text(slide, 5.75, y, 6.75, 0.40, url, 9, SLATE)
            # Make the title clickable while keeping it editable.
            p = shape.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].hyperlink.address = url
            y += 0.58
        add_footer(slide, page, "点击资料标题可打开官方页面")
        page += 1

    out = ROOT / "ai-database-lakehouse-report-editable.pptx"
    prs.save(out)
    print(f"generated {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
